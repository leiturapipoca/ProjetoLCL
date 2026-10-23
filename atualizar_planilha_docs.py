import argparse
import re
import unicodedata
from pathlib import Path
import pandas as pd
import openpyxl
from docx import Document

from config import ONEDRIVE_BASE_DIR, inicializar_pastas
from controle_downloads import chave_projeto

# ==============================================================================
# 1. CONFIGURAÇÕES DE CAMINHOS PADRÃO
# ==============================================================================
PASTA_NOTAS_PADRAO = Path(r"C:\automacao_opee\pasta_notas_tecnicas")
PLANILHA_EXCEL_PADRAO = Path(r"C:\automacao_opee\planilha_empresas.xlsx")
ARQUIVO_WARNINGS_PADRAO = ONEDRIVE_BASE_DIR / "warnings_notas_tecnicas.txt"

ANO_PADRAO = "2026"

# Regex separadas: uma estrita para PE e outra para detectar projetos que devem ser ignorados (ex: PG, PD)
REGEX_PROJETO_PE = re.compile(r"\bPE-\d+-\d+\s*/\s*\d{4}\b", re.IGNORECASE)
REGEX_PROJETO_OUTROS = re.compile(r"\b(?!PE-)[A-Z]{2,3}-\d+-\d+\s*/\s*\d{4}\b", re.IGNORECASE)

# Regex para "NOTA TÉCNICA Nº XX/2026"
REGEX_NT_CABECALHO = re.compile(
    r"NOTA\s+T[ÉE]CNICA\s*(?:N[º°o.]*)?\s*([0-9X]+)\s*/\s*(\d{4})",
    re.IGNORECASE,
)

# Regex para o primeiro número entre "_" e "_" no nome do arquivo (_XXX_)
REGEX_NUMERO_ARQUIVO = re.compile(r"_\s*(\d+)\s*_")


def normalizar_nome_coluna(texto: str) -> str:
    if not texto:
        return ""
    sem_acento = "".join(
        c for c in unicodedata.normalize("NFD", str(texto))
        if unicodedata.category(c) != "Mn"
    )
    return re.sub(r"[^a-zA-Z0-9]+", " ", sem_acento).strip().lower()


def encontrar_coluna(colunas_planilha: list[str], nomes_possiveis: list[str]) -> str:
    mapa = {normalizar_nome_coluna(c): c for c in colunas_planilha}
    for nome in nomes_possiveis:
        if nome in colunas_planilha:
            return nome
        norm = normalizar_nome_coluna(nome)
        if norm in mapa:
            return mapa[norm]
    raise ValueError(f"Coluna não encontrada. Esperado um de {nomes_possiveis} em {colunas_planilha}")


def campo_esta_vazio(valor) -> bool:
    if pd.isna(valor) or valor is None:
        return True
    texto = str(valor).strip()
    return not bool(texto) or texto.lower() in ("nan", "none", "-", "n/a", "não", "nao")


def limpar_texto_celula(texto_celula: str) -> str:
    if not texto_celula:
        return ""
    return re.sub(r"\s+", " ", texto_celula).strip()


def eh_zero_reais(texto_celula: str) -> bool:
    """Retorna True se o valor na coluna de glosa for R$ 0,00 (ou vazio)."""
    if not texto_celula or not texto_celula.strip():
        return True
    limpo = re.sub(r"\s+", "", texto_celula).upper()
    match_num = re.search(r"(\d[\d.,]*)", limpo)
    if not match_num:
        return True
    num_str = match_num.group(1).replace(".", "").replace(",", ".")
    try:
        return abs(float(num_str)) < 0.005
    except ValueError:
        return limpo in ("R$0,00", "0,00", "0", "-")


def extrair_textos_documento(doc: Document) -> tuple[str, str]:
    partes_topo = []
    for section in doc.sections:
        for p in section.header.paragraphs:
            if p.text.strip():
                partes_topo.append(p.text.strip())
        for t in section.header.tables:
            for row in t.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        partes_topo.append(cell.text.strip())

    paragrafos_corpo = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    textos_primeira_tabela = []
    if doc.tables:
        for row in doc.tables[0].rows[:3]:
            for cell in row.cells:
                if cell.text.strip():
                    textos_primeira_tabela.append(cell.text.strip())

    texto_topo = "\n".join(partes_topo + textos_primeira_tabela + paragrafos_corpo[:15])
    texto_corpo_completo = "\n".join(paragrafos_corpo)
    return texto_topo, texto_corpo_completo


def extrair_codigo_projeto(doc: Document, texto_corpo: str) -> tuple[str | None, str | None]:
    """
    Retorna (codigo_pe, codigo_ignorado).
    Prioriza o bloco 'I - DO OBJETIVO'. Se encontrar apenas projeto não-PE (ex: PG-...),
    retorna (None, codigo_outro) para que o arquivo seja ignorado sem gerar warning.
    """
    match_bloco = re.search(
        r"(?:I\s*[-–—.]?\s*DO\s+OBJETIVO)(.*?)(?:II\s*[-–—.]|\n\s*2\s*[-–—.]|$)",
        texto_corpo,
        re.IGNORECASE | re.DOTALL,
    )
    if match_bloco:
        trecho_objetivo = match_bloco.group(1)
        codigos_pe = REGEX_PROJETO_PE.findall(trecho_objetivo)
        if codigos_pe:
            return chave_projeto(codigos_pe[-1]), None
        codigos_outros = REGEX_PROJETO_OUTROS.findall(trecho_objetivo)
        if codigos_outros:
            return None, codigos_outros[-1].upper()

    # Fallback: procura no corpo inteiro
    codigos_pe_corpo = REGEX_PROJETO_PE.findall(texto_corpo)
    if codigos_pe_corpo:
        return chave_projeto(codigos_pe_corpo[0]), None

    codigos_outros_corpo = REGEX_PROJETO_OUTROS.findall(texto_corpo)
    if codigos_outros_corpo:
        return None, codigos_outros_corpo[0].upper()

    # Fallback extra: procura nas tabelas
    for tabela in doc.tables:
        for row in tabela.rows:
            for cell in row.cells:
                achados_pe = REGEX_PROJETO_PE.findall(cell.text)
                if achados_pe:
                    return chave_projeto(achados_pe[0]), None
                achados_outros = REGEX_PROJETO_OUTROS.findall(cell.text)
                if achados_outros:
                    return None, achados_outros[0].upper()

    return None, None


def extrair_nota_tecnica(texto_topo: str, texto_corpo: str, nome_arquivo: str) -> tuple[str | None, bool]:
    match_doc = REGEX_NT_CABECALHO.search(texto_topo) or REGEX_NT_CABECALHO.search(texto_corpo)

    esqueceu_xx = False
    if match_doc:
        num_bruto = match_doc.group(1).strip().upper()
        ano_doc = match_doc.group(2).strip()
        if "X" not in num_bruto and num_bruto.isdigit():
            return f"{int(num_bruto)}/{ano_doc}", False
        else:
            esqueceu_xx = True

    stem = Path(nome_arquivo).stem
    match_arq = REGEX_NUMERO_ARQUIVO.search(stem)
    if match_arq:
        num_arq = int(match_arq.group(1))
        return f"{num_arq}/{ANO_PADRAO}", esqueceu_xx

    return None, esqueceu_xx


def extrair_valor_por_total_docx(doc: Document) -> tuple[str | None, bool, str]:
    """
    Procura nas tabelas a célula contendo 'Total' (ex: Total | valor | glosado).
    - O Valor Aprovado está na coluna imediatamente após 'Total'.
    - Se houver mais uma coluna à direita ('glosado') e ela não for R$ 0,00,
      marca warning como 'Valor glosado'.
    """
    for tabela in doc.tables:
        for row in tabela.rows:
            # Desduplica células mescladas horizontalmente usando o elemento XML (_tc)
            celulas_unicas = []
            tc_vistos = set()
            for cell in row.cells:
                if id(cell._tc) not in tc_vistos:
                    tc_vistos.add(id(cell._tc))
                    txt = limpar_texto_celula(cell.text)
                    if txt:
                        celulas_unicas.append(txt)

            for idx, texto_cel in enumerate(celulas_unicas):
                # Verifica se a célula contém a palavra "Total" (ignorando "Subtotal")
                if re.search(r"\bTotal\b", texto_cel, re.IGNORECASE):
                    if idx + 1 < len(celulas_unicas):
                        valor_aprovado = celulas_unicas[idx + 1]
                        # Confirma que a célula seguinte contém dígitos
                        if not re.search(r"\d", valor_aprovado):
                            continue

                        texto_glosa = ""
                        teve_glosa = False
                        if idx + 2 < len(celulas_unicas):
                            texto_glosa = celulas_unicas[idx + 2]
                            if re.search(r"\d", texto_glosa) and not eh_zero_reais(texto_glosa):
                                teve_glosa = True

                        return valor_aprovado, teve_glosa, texto_glosa

    return None, False, ""


def analisar_docx(caminho_docx: Path) -> dict:
    doc = Document(str(caminho_docx))
    texto_topo, texto_corpo = extrair_textos_documento(doc)

    codigo_pe, codigo_ignorado = extrair_codigo_projeto(doc, texto_corpo)
    if codigo_ignorado and not codigo_pe:
        return {
            "arquivo": caminho_docx.name,
            "ignorado_nao_pe": True,
            "codigo_ignorado": codigo_ignorado,
        }

    nota_tecnica, esqueceu_xx = extrair_nota_tecnica(texto_topo, texto_corpo, caminho_docx.name)
    valor_aprovado, teve_glosa, texto_glosa = extrair_valor_por_total_docx(doc)

    return {
        "arquivo": caminho_docx.name,
        "ignorado_nao_pe": False,
        "codigo_projeto": codigo_pe,
        "nota_tecnica": nota_tecnica,
        "valor_aprovado": valor_aprovado,
        "esqueceu_xx": esqueceu_xx,
        "teve_glosa": teve_glosa,
        "texto_glosa": texto_glosa,
    }


def executar_atualizacao_planilha(
    pasta_notas: Path,
    caminho_planilha: Path,
    caminho_warnings: Path,
    nome_aba: str | int = 0,
    recursivo: bool = False,
):
    inicializar_pastas()

    if not pasta_notas.exists():
        raise FileNotFoundError(f"Pasta de Notas Técnicas não encontrada: {pasta_notas}")
    if not caminho_planilha.exists():
        raise FileNotFoundError(f"Planilha não encontrada: {caminho_planilha}")

    print("=" * 75)
    print(" LENDO ARQUIVOS .DOCX (APENAS PE) E ATUALIZANDO PLANILHA")
    print("=" * 75)

    iterador = pasta_notas.rglob("*") if recursivo else pasta_notas.iterdir()
    arquivos_docx = sorted(
        arq for arq in iterador
        if arq.is_file() and arq.suffix.lower() == ".docx" and not arq.name.startswith("~$")
    )

    print(f"[1/4] Encontrados {len(arquivos_docx)} arquivo(s) .docx na pasta.")

    dados_extraidos: dict[str, dict] = {}
    warnings: list[str] = []
    ignorados_nao_pe = 0

    for idx, arq in enumerate(arquivos_docx, start=1):
        try:
            info = analisar_docx(arq)
        except Exception as erro:
            warnings.append(f"[ERRO DE LEITURA] Arquivo: {arq.name} | Erro: {erro}")
            continue

        if info.get("ignorado_nao_pe"):
            ignorados_nao_pe += 1
            print(f"   [{idx}/{len(arquivos_docx)}] [IGNORADO - NÃO É PE] {info['codigo_ignorado']} ({arq.name})")
            continue

        cod = info["codigo_projeto"]
        nt = info["nota_tecnica"]
        val = info["valor_aprovado"]

        if info["esqueceu_xx"]:
            warnings.append(
                f"-Esqueceu de mudar na própria nota técnica- | Projeto: {cod or 'N/A'} | "
                f"NT atribuída pelo nome: {nt or 'N/A'} | Arquivo: {arq.name}"
            )

        if info["teve_glosa"]:
            warnings.append(
                f"Valor glosado | Projeto: {cod or 'N/A'} | NT: {nt or 'N/A'} | "
                f"Valor glosado: {info['texto_glosa']} | Valor aprovado (após Total): {val} | Arquivo: {arq.name}"
            )

        if not cod:
            warnings.append(f"[PROJETO PE NÃO IDENTIFICADO] Regex PE-XXXXX-XXXX/XXXX não encontrada em: {arq.name}")
            continue

        dados_extraidos[cod] = info
        print(f"   [{idx}/{len(arquivos_docx)}] {cod} -> NT: {nt} | Valor: {val}")

    # 2. Salvar arquivo de warnings
    caminho_warnings.parent.mkdir(parents=True, exist_ok=True)
    with open(caminho_warnings, "w", encoding="utf-8") as f:
        if warnings:
            f.write("\n".join(warnings) + "\n")
    print(f"\n[2/4] Arquivo de warnings salvo em '{caminho_warnings.name}' ({len(warnings)} ocorrência(s)).")

    # 3. Ler planilha e identificar atualizações (ignorando projetos que não sejam PE)
    print(f"[3/4] Analisando planilha '{caminho_planilha.name}'...")
    df = pd.read_excel(caminho_planilha, sheet_name=nome_aba)
    colunas_originais = [str(c).strip() for c in df.columns]
    df.columns = colunas_originais

    col_cod = encontrar_coluna(colunas_originais, ["Codigo do Projeto", "Código do Projeto"])
    col_aneel = encontrar_coluna(colunas_originais, ["Avaliado Pela ANEEL"])
    col_nt = encontrar_coluna(colunas_originais, ["Nota Técnica", "Nota Tecnica"])
    col_valor = encontrar_coluna(colunas_originais, ["Valor Aprovado"])

    idx_col_aneel = colunas_originais.index(col_aneel) + 1
    idx_col_nt = colunas_originais.index(col_nt) + 1
    idx_col_valor = colunas_originais.index(col_valor) + 1

    amostra_aneel = df[col_aneel].dropna().astype(str).str.strip().tolist()
    texto_sim = "SIM" if "SIM" in amostra_aneel else "Sim"

    atualizacoes = []

    for idx_df, row in df.iterrows():
        cod_bruto = str(row[col_cod]).strip() if pd.notna(row[col_cod]) else ""
        if not cod_bruto:
            continue

        try:
            chave = chave_projeto(cod_bruto)
        except ValueError:
            continue

        # Ignora se não começar com PE-
        if not chave.upper().startswith("PE-"):
            continue

        if chave not in dados_extraidos:
            continue

        info_doc = dados_extraidos[chave]
        nt_doc = info_doc["nota_tecnica"]
        val_doc = info_doc["valor_aprovado"]

        aneel_atual = str(row[col_aneel]).strip().upper() if pd.notna(row[col_aneel]) else ""
        nt_vazia = campo_esta_vazio(row[col_nt])
        valor_vazio = campo_esta_vazio(row[col_valor])

        era_nao = aneel_atual in ("NÃO", "NAO", "N", "")
        era_sim_incompleto = aneel_atual in ("SIM", "S") and (nt_vazia or valor_vazio)

        if era_nao or era_sim_incompleto:
            linha_excel = idx_df + 2
            novo_aneel = texto_sim if era_nao else None
            nova_nt = nt_doc if (era_nao or nt_vazia) and nt_doc else None
            novo_val = val_doc if (era_nao or valor_vazio) and val_doc else None
            atualizacoes.append((linha_excel, chave, novo_aneel, nova_nt, novo_val))

    # 4. Gravar alterações na planilha Excel
    print(f"[4/4] Aplicando atualizações em {len(atualizacoes)} linha(s) da planilha...")
    if atualizacoes:
        while True:
            try:
                wb = openpyxl.load_workbook(caminho_planilha)
                ws = wb[nome_aba] if isinstance(nome_aba, str) else wb.worksheets[int(nome_aba)]

                for linha_excel, chave, novo_aneel, nova_nt, novo_val in atualizacoes:
                    if novo_aneel is not None:
                        ws.cell(row=linha_excel, column=idx_col_aneel, value=novo_aneel)
                    if nova_nt is not None:
                        ws.cell(row=linha_excel, column=idx_col_nt, value=nova_nt)
                    if novo_val is not None:
                        ws.cell(row=linha_excel, column=idx_col_valor, value=novo_val)

                wb.save(caminho_planilha)
                wb.close()
                break
            except PermissionError:
                print("\n [!] ATENÇÃO: A planilha está aberta no Excel!")
                input(" >>> Feche a planilha no Excel e pressione ENTER para salvar... ")

    print("\n" + "=" * 75)
    print(" RESUMO FINAL (LEITURA DE DOCX)")
    print("=" * 75)
    print(f" - Projetos PE válidos extraídos                  : {len(dados_extraidos)}")
    print(f" - Projetos ignorados (PG, PD, etc.)              : {ignorados_nao_pe}")
    print(f" - Projetos atualizados na planilha Excel         : {len(atualizacoes)}")
    print(f" - Ocorrências salvas no arquivo de warnings      : {len(warnings)}")
    print("=" * 75)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Extrai Projeto (apenas PE), Nota Técnica e Valor (após Total) dos arquivos .docx e atualiza a planilha."
    )
    parser.add_argument("--pasta", "-p", type=Path, default=PASTA_NOTAS_PADRAO, help="Pasta com as Notas Técnicas.")
    parser.add_argument("--planilha", "-x", type=Path, default=PLANILHA_EXCEL_PADRAO, help="Planilha Excel (.xlsx).")
    parser.add_argument("--warnings", "-w", type=Path, default=ARQUIVO_WARNINGS_PADRAO, help="Arquivo TXT de warnings.")
    parser.add_argument("--aba", "-a", default=0, help="Nome ou número da aba da planilha (padrão: 0).")
    parser.add_argument("--recursivo", "-r", action="store_true", help="Busca .docx também em subpastas.")

    args = parser.parse_args()
    executar_atualizacao_planilha(
        pasta_notas=args.pasta,
        caminho_planilha=args.planilha,
        caminho_warnings=args.warnings,
        nome_aba=args.aba,
        recursivo=args.recursivo,
    )
