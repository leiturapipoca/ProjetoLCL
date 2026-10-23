import argparse
import re
import unicodedata
from pathlib import Path
import pandas as pd
import openpyxl
import pdfplumber

from config import ONEDRIVE_BASE_DIR, inicializar_pastas
from controle_downloads import chave_projeto

# ==============================================================================
# 1. CONFIGURAÇÕES DE CAMINHOS PADRÃO
# ==============================================================================
PASTA_NOTAS_PADRAO = Path(r"C:\automacao_opee\pasta_notas_tecnicas")
PLANILHA_EXCEL_PADRAO = Path(r"C:\automacao_opee\planilha_empresas.xlsx")
ARQUIVO_WARNINGS_PADRAO = ONEDRIVE_BASE_DIR / "warnings_notas_tecnicas.txt"

ANO_PADRAO = "2026"

# Regex para Código do Projeto (PE-XXXXX-XXXX/XXXX, aceitando também PD/PG)
REGEX_PROJETO = re.compile(r"\b(?:PE|PD|PG)-\d+-\d+\s*/\s*\d{4}\b", re.IGNORECASE)

# Regex para "NOTA TÉCNICA Nº XX/2026" logo após a logo no PDF
REGEX_NT_CABECALHO = re.compile(
    r"NOTA\s+T[ÉE]CNICA\s*(?:N[º°o.]*)?\s*([0-9X]+)\s*/\s*(\d{4})",
    re.IGNORECASE,
)

# Regex para o primeiro número entre "_" e "_" no nome do arquivo (_XXX_)
REGEX_NUMERO_ARQUIVO = re.compile(r"_\s*(\d+)\s*_")


def normalizar_nome_coluna(texto: str) -> str:
    """Remove acentos e espaços extras para localizar colunas na planilha."""
    if not texto:
        return ""
    sem_acento = "".join(
        c for c in unicodedata.normalize("NFD", str(texto))
        if unicodedata.category(c) != "Mn"
    )
    return re.sub(r"[^a-zA-Z0-9]+", " ", sem_acento).strip().lower()


def encontrar_coluna(colunas_planilha: list[str], nomes_possiveis: list[str]) -> str:
    """Retorna o nome exato da coluna na planilha."""
    mapa = {normalizar_nome_coluna(c): c for c in colunas_planilha}
    for nome in nomes_possiveis:
        if nome in colunas_planilha:
            return nome
        norm = normalizar_nome_coluna(nome)
        if norm in mapa:
            return mapa[norm]
    raise ValueError(f"Coluna não encontrada. Esperado um de {nomes_possiveis} em {colunas_planilha}")


def campo_esta_vazio(valor) -> bool:
    """Verifica se uma célula da planilha está vazia ou sem preenchimento válido."""
    if pd.isna(valor) or valor is None:
        return True
    texto = str(valor).strip()
    return not bool(texto) or texto.lower() in ("nan", "none", "-", "n/a", "não", "nao")


def limpar_celula_pdf(valor_celula: str | None) -> str:
    """Remove quebras de linha internas da célula do PDF e espaços duplicados."""
    if not valor_celula:
        return ""
    return re.sub(r"\s+", " ", str(valor_celula)).strip()


def eh_zero_reais(texto_celula: str) -> bool:
    """
    Verifica se o valor na posição 4x3 equivale a R$ 0,00.
    Retorna False se houver qualquer valor glosado diferente de zero.
    """
    if not texto_celula or not texto_celula.strip():
        return True
    limpo = re.sub(r"\s+", "", texto_celula).upper()
    match_num = re.search(r"(\d[\d.,]*)", limpo)
    if not match_num:
        return True
    num_str = match_num.group(1).replace(".", "").replace(",", ".")
    try:
        return abs(float(num_str)) < 0.005
    except ValueError:
        return limpo in ("R$0,00", "0,00", "0", "-")


def extrair_codigo_projeto(texto_completo: str) -> str | None:
    """
    Procura o código PE-XXXXX-XXXX/XXXX priorizando a última ocorrência
    dentro do bloco 'I - DO OBJETIVO'.
    """
    match_bloco = re.search(
        r"(?:I\s*[-–—.]?\s*DO\s+OBJETIVO)(.*?)(?:II\s*[-–—.]|\n\s*2\s*[-–—.]|$)",
        texto_completo,
        re.IGNORECASE | re.DOTALL,
    )
    if match_bloco:
        trecho_objetivo = match_bloco.group(1)
        codigos_no_objetivo = REGEX_PROJETO.findall(trecho_objetivo)
        if codigos_no_objetivo:
            # Pega a última ocorrência escrita na descrição do bloco I - DO OBJETIVO
            return chave_projeto(codigos_no_objetivo[-1])

    # Fallback: primeira ocorrência no texto geral do PDF
    codigos_geral = REGEX_PROJETO.findall(texto_completo)
    if codigos_geral:
        return chave_projeto(codigos_geral[0])

    return None


def extrair_nota_tecnica(texto_primeira_pagina: str, texto_completo: str, nome_arquivo: str) -> tuple[str | None, bool]:
    """
    1ª tentativa: 'NOTA TÉCNICA Nº XX/2026' logo após a logo na 1ª página.
    2ª tentativa: se estiver preenchido com 'XX' (esqueceram de mudar), pega o primeiro
                  número entre '_ _' no nome do arquivo (_XXX_) e sinaliza warning.
    """
    match_pdf = REGEX_NT_CABECALHO.search(texto_primeira_pagina) or REGEX_NT_CABECALHO.search(texto_completo)

    esqueceu_xx = False
    if match_pdf:
        num_bruto = match_pdf.group(1).strip().upper()
        ano_pdf = match_pdf.group(2).strip()
        if "X" not in num_bruto and num_bruto.isdigit():
            return f"{int(num_bruto)}/{ano_pdf}", False
        else:
            # Documento está com 'XX/2026'
            esqueceu_xx = True

    # Fallback pelo nome do arquivo (_XXX_)
    stem = Path(nome_arquivo).stem
    match_arq = REGEX_NUMERO_ARQUIVO.search(stem)
    if match_arq:
        num_arq = int(match_arq.group(1))
        return f"{num_arq}/{ANO_PADRAO}", esqueceu_xx

    return None, esqueceu_xx


def extrair_valor_das_tabelas_pdf(tabelas_brutas: list[list[list[str | None]]]) -> tuple[str | None, bool, str]:
    """
    Regra das Tabelas no PDF:
    - Filtra tabelas que tenham pelo menos 4 linhas e 3 colunas.
    - Se for uma tabela 4x3 (4 linhas x 3 colunas):
        * Puxa o valor de 4x2 (linha índice 3, coluna índice 1).
        * Checa 4x3 (linha índice 3, coluna índice 2). Se != R$ 0,00 -> warning 'Valor glosado'.
    - Em tabelas maiores (mais de 4 linhas ou mais de 3 colunas):
        * Puxa o valor sempre da posição 4x3 (linha índice 3, coluna índice 2).
    """
    tabelas_candidatas = []

    for tab in tabelas_brutas:
        # Remove linhas inteiramente vazias caso o PDF tenha bordas fantasma
        linhas_validas = [
            [limpar_celula_pdf(cel) for cel in linha]
            for linha in tab
            if linha and any(limpar_celula_pdf(c) for c in linha)
        ]
        if len(linhas_validas) >= 4 and len(linhas_validas[3]) >= 3:
            tabelas_candidatas.append(linhas_validas)

    if not tabelas_candidatas:
        return None, False, ""

    # Se houver mais de uma tabela >= 4x3, escolhe a que contém dígitos na 4ª linha
    tabela_alvo = tabelas_candidatas[0]
    if len(tabelas_candidatas) > 1:
        for t in tabelas_candidatas:
            texto_linha_4 = " ".join(t[3])
            if re.search(r"\d", texto_linha_4):
                tabela_alvo = t
                break

    qtd_linhas = len(tabela_alvo)
    qtd_colunas = max(len(linha) for linha in tabela_alvo)

    # Caso 1: Tabela única 4x3
    if len(tabelas_candidatas) == 1 and qtd_linhas == 4 and qtd_colunas == 3:
        valor_4x2 = tabela_alvo[3][1]
        valor_4x3 = tabela_alvo[3][2]

        teve_glosa = not eh_zero_reais(valor_4x3)
        return valor_4x2, teve_glosa, valor_4x3

    # Caso 2: Tabelas maiores -> posição 4x3 (4ª linha, 3ª coluna)
    valor_4x3 = tabela_alvo[3][2]
    return valor_4x3, False, ""


def analisar_pdf(caminho_pdf: Path) -> dict:
    """Abre um arquivo .pdf com pdfplumber e extrai Projeto, Nota Técnica, Valor e Warnings."""
    textos_paginas = []
    todas_tabelas = []

    with pdfplumber.open(str(caminho_pdf)) as pdf:
        for pagina in pdf.pages:
            txt = pagina.extract_text() or ""
            if txt.strip():
                textos_paginas.append(txt)

            tabelas_pag = pagina.extract_tables() or []
            if tabelas_pag:
                todas_tabelas.extend(tabelas_pag)

    texto_primeira_pagina = textos_paginas[0] if textos_paginas else ""
    texto_completo = "\n".join(textos_paginas)

    codigo_projeto = extrair_codigo_projeto(texto_completo)
    nota_tecnica, esqueceu_xx = extrair_nota_tecnica(texto_primeira_pagina, texto_completo, caminho_pdf.name)
    valor_aprovado, teve_glosa, texto_glosa = extrair_valor_das_tabelas_pdf(todas_tabelas)

    return {
        "arquivo": caminho_pdf.name,
        "codigo_projeto": codigo_projeto,
        "nota_tecnica": nota_tecnica,
        "valor_aprovado": valor_aprovado,
        "esqueceu_xx": esqueceu_xx,
        "teve_glosa": teve_glosa,
        "texto_glosa": texto_glosa,
    }


def executar_atualizacao_pdfs(
    pasta_notas: Path,
    caminho_planilha: Path,
    caminho_warnings: Path,
    nome_aba: str | int = 0,
    recursivo: bool = False,
):
    inicializar_pastas()

    if not pasta_notas.exists():
        raise FileNotFoundError(f"Pasta de Notas Técnicas não encontrada: {pasta_notas}")
    if not caminho_planilha.exists():
        raise FileNotFoundError(f"Planilha não encontrada: {caminho_planilha}")

    print("=" * 75)
    print(" LENDO ARQUIVOS .PDF E ATUALIZANDO PLANILHA DE PROJETOS")
    print("=" * 75)

    # Filtra exclusivamente os arquivos .pdf da pasta
    iterador = pasta_notas.rglob("*") if recursivo else pasta_notas.iterdir()
    arquivos_pdf = sorted(
        arq for arq in iterador
        if arq.is_file() and arq.suffix.lower() == ".pdf"
    )

    print(f"[1/4] Encontrados {len(arquivos_pdf)} arquivo(s) .pdf na pasta.")

    dados_extraidos: dict[str, dict] = {}
    warnings: list[str] = []

    for idx, arq in enumerate(arquivos_pdf, start=1):
        try:
            info = analisar_pdf(arq)
        except Exception as erro:
            warnings.append(f"[ERRO AO LER PDF] Arquivo: {arq.name} | Erro: {erro}")
            continue

        cod = info["codigo_projeto"]
        nt = info["nota_tecnica"]
        val = info["valor_aprovado"]

        if info["esqueceu_xx"]:
            warnings.append(
                f"-Esqueceu de mudar na própria nota técnica- | Projeto: {cod or 'N/A'} | "
                f"NT atribuída pelo nome: {nt or 'N/A'} | Arquivo: {arq.name}"
            )

        if info["teve_glosa"]:
            warnings.append(
                f"Valor glosado | Projeto: {cod or 'N/A'} | NT: {nt or 'N/A'} | "
                f"Valor em 4x3 (Glosa): {info['texto_glosa']} | Valor puxado (4x2): {val} | Arquivo: {arq.name}"
            )

        if not cod:
            warnings.append(f"[PROJETO NÃO IDENTIFICADO] Regex PE-XXXXX-XXXX/XXXX não encontrada em: {arq.name}")
            continue

        dados_extraidos[cod] = info
        print(f"   [{idx}/{len(arquivos_pdf)}] {cod} -> NT: {nt} | Valor: {val}")

    # 2. Salvar arquivo de warnings
    caminho_warnings.parent.mkdir(parents=True, exist_ok=True)
    with open(caminho_warnings, "w", encoding="utf-8") as f:
        if warnings:
            f.write("\n".join(warnings) + "\n")
    print(f"\n[2/4] Arquivo de warnings salvo em '{caminho_warnings.name}' ({len(warnings)} ocorrência(s)).")

    # 3. Ler planilha e verificar quais linhas precisam de atualização
    print(f"[3/4] Analisando planilha '{caminho_planilha.name}'...")
    df = pd.read_excel(caminho_planilha, sheet_name=nome_aba)
    colunas_originais = [str(c).strip() for c in df.columns]
    df.columns = colunas_originais

    col_cod = encontrar_coluna(colunas_originais, ["Codigo do Projeto", "Código do Projeto"])
    col_aneel = encontrar_coluna(colunas_originais, ["Avaliado Pela ANEEL"])
    col_nt = encontrar_coluna(colunas_originais, ["Nota Técnica", "Nota Tecnica"])
    col_valor = encontrar_coluna(colunas_originais, ["Valor Aprovado"])

    idx_col_aneel = colunas_originais.index(col_aneel) + 1
    idx_col_nt = colunas_originais.index(col_nt) + 1
    idx_col_valor = colunas_originais.index(col_valor) + 1

    # Preserva o padrão de maiúsculas/minúsculas do "SIM" usado na planilha
    amostra_aneel = df[col_aneel].dropna().astype(str).str.strip().tolist()
    texto_sim = "SIM" if "SIM" in amostra_aneel else "Sim"

    atualizacoes = []  # (linha_excel, codigo, novo_aneel, nova_nt, novo_valor)

    for idx_df, row in df.iterrows():
        cod_bruto = str(row[col_cod]).strip() if pd.notna(row[col_cod]) else ""
        if not cod_bruto:
            continue

        try:
            chave = chave_projeto(cod_bruto)
        except ValueError:
            continue

        if chave not in dados_extraidos:
            continue

        info_pdf = dados_extraidos[chave]
        nt_pdf = info_pdf["nota_tecnica"]
        val_pdf = info_pdf["valor_aprovado"]

        aneel_atual = str(row[col_aneel]).strip().upper() if pd.notna(row[col_aneel]) else ""
        nt_vazia = campo_esta_vazio(row[col_nt])
        valor_vazio = campo_esta_vazio(row[col_valor])

        era_nao = aneel_atual in ("NÃO", "NAO", "N", "")
        era_sim_incompleto = aneel_atual in ("SIM", "S") and (nt_vazia or valor_vazio)

        if era_nao or era_sim_incompleto:
            linha_excel = idx_df + 2
            novo_aneel = texto_sim if era_nao else None
            nova_nt = nt_pdf if (era_nao or nt_vazia) and nt_pdf else None
            novo_val = val_pdf if (era_nao or valor_vazio) and val_pdf else None
            atualizacoes.append((linha_excel, chave, novo_aneel, nova_nt, novo_val))

    # 4. Gravar alterações na planilha Excel preservando a formatação
    print(f"[4/4] Aplicando atualizações em {len(atualizacoes)} linha(s) da planilha...")
    if atualizacoes:
        while True:
            try:
                wb = openpyxl.load_workbook(caminho_planilha)
                ws = wb[nome_aba] if isinstance(nome_aba, str) else wb.worksheets[int(nome_aba)]

                for linha_excel, chave, novo_aneel, nova_nt, novo_val in atualizacoes:
                    if novo_aneel is not None:
                        ws.cell(row=linha_excel, column=idx_col_aneel, value=novo_aneel)
                    if nova_nt is not None:
                        ws.cell(row=linha_excel, column=idx_col_nt, value=nova_nt)
                    if novo_val is not None:
                        ws.cell(row=linha_excel, column=idx_col_valor, value=novo_val)

                wb.save(caminho_planilha)
                wb.close()
                break
            except PermissionError:
                print("\n [!] ATENÇÃO: A planilha está aberta no Excel!")
                input(" >>> Feche a planilha no Excel e pressione ENTER para salvar... ")

    print("\n" + "=" * 75)
    print(" RESUMO FINAL (LEITURA DE PDFS)")
    print("=" * 75)
    print(f" - Documentos .pdf processados com código válido  : {len(dados_extraidos)}")
    print(f" - Projetos atualizados na planilha Excel         : {len(atualizacoes)}")
    print(f" - Ocorrências salvas no arquivo de warnings      : {len(warnings)}")
    print(f" - Caminho do arquivo de warnings                 : {caminho_warnings}")
    print("=" * 75)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Extrai Projeto, Nota Técnica e Valor diretamente dos arquivos .pdf e atualiza a planilha."
    )
    parser.add_argument("--pasta", "-p", type=Path, default=PASTA_NOTAS_PADRAO, help="Pasta com as Notas Técnicas (.pdf).")
    parser.add_argument("--planilha", "-x", type=Path, default=PLANILHA_EXCEL_PADRAO, help="Planilha Excel (.xlsx).")
    parser.add_argument("--warnings", "-w", type=Path, default=ARQUIVO_WARNINGS_PADRAO, help="Arquivo TXT de warnings.")
    parser.add_argument("--aba", "-a", default=0, help="Nome ou número da aba da planilha (padrão: 0).")
    parser.add_argument("--recursivo", "-r", action="store_true", help="Busca .pdf também em subpastas.")

    args = parser.parse_args()
    executar_atualizacao_pdfs(
        pasta_notas=args.pasta,
        caminho_planilha=args.planilha,
        caminho_warnings=args.warnings,
        nome_aba=args.aba,
        recursivo=args.recursivo,
    )
